#!/usr/bin/env python3
"""Parse one weekly TRM commission consolidation statement workbook.

The consolidation statements are .xlsx/.xlsm workbooks with, on one sheet:
  - a "Total:" (older files: "Statement Value:") cell whose neighbour holds the
    statement total,
  - a "Date:" cell whose neighbour holds the statement date,
  - a table beginning at an "Adviser" header row, with columns
    Adviser | Date | Lenders | Policy reference | Type | First name | Surname |
    Class | Commission.

Public API:
    parse_statement(path, filename=None) -> dict
        On success: {file, date, total, items:[{date,lender,ref,type,firstName,
                     surname,class,amount}], [sumMismatch, itemsSum],
                     [totalDerivedFromItems]}
        On a workbook that isn't a statement: {file, skipped: "<reason>"}

This module contains NO data — it is pure parsing logic, used by both the local
Google refresh (scripts/gdrive_refresh.py) and any manual extraction.
"""
import datetime
import os
import re

from openpyxl import load_workbook

MONTHS = {m.lower(): i for i, m in enumerate(
    ["Jan", "Feb", "Mar", "Apr", "May", "Jun",
     "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"], 1)}


def norm_amount(value):
    """Money cell -> float or None. Handles '£1,407.81', '-£194.97', '(x)'."""
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return float(value)
    s = str(value).strip().replace("£", "").replace(",", "").replace(" ", "")
    if s.startswith("(") and s.endswith(")"):
        s = "-" + s[1:-1]
    if s in ("", "-"):
        return None
    try:
        return float(s)
    except ValueError:
        return None


def norm_date(value):
    """Date cell -> 'YYYY-MM-DD' or None. Handles datetimes and '23-Jan-26'."""
    if value is None:
        return None
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    s = str(value).strip()
    m = re.match(r"^(\d{1,2})[-/ ]([A-Za-z]{3,9})[-/ ](\d{2,4})$", s)
    if m:
        d, mon, y = m.groups()
        mi = MONTHS.get(mon[:3].lower())
        if mi:
            y = int(y)
            y = y + 2000 if y < 100 else y
            return "%04d-%02d-%02d" % (y, mi, int(d))
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})$", s)
    if m:
        d, mo, y = (int(x) for x in m.groups())
        y = y + 2000 if y < 100 else y
        return "%04d-%02d-%02d" % (y, mo, d)
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return m.group(0)
    return None


def date_from_filename(filename):
    """'...30-01-2026.xlsx' or '...11.11.2024.xlsx' -> 'YYYY-MM-DD' or None."""
    m = re.search(r"(\d{2})[-.](\d{2})[-.](\d{4})", filename or "")
    if m:
        d, mo, y = m.groups()
        return "%s-%s-%s" % (y, mo, d)
    return None


def _cell_text(value):
    return str(value).strip() if value is not None else ""


_HEADER_MAP = {
    "adviser": "adviser", "date": "date", "lenders": "lender", "lender": "lender",
    "policy reference": "ref", "policy ref": "ref", "reference": "ref",
    "type": "type", "first name": "firstName", "firstname": "firstName",
    "surname": "surname", "class": "class", "commission": "amount", "amount": "amount",
}


def parse_statement(path, filename=None):
    """Parse a single statement workbook. See module docstring for the shape."""
    filename = filename or os.path.basename(path)
    wb = load_workbook(path, data_only=True, read_only=True)
    parsed = None
    try:
        for ws in wb.worksheets:
            rows = [row for row in ws.iter_rows(values_only=True)]
            total_val = None
            date_val = None
            header_idx = None
            header_row = None
            for i, row in enumerate(rows):
                for j, c in enumerate(row):
                    t = _cell_text(c).lower()
                    if header_idx is None and total_val is None and \
                            t in ("total:", "total", "statement value:", "statement value"):
                        for k in range(j + 1, min(j + 6, len(row))):
                            a = norm_amount(row[k])
                            if a is not None:
                                total_val = a
                                break
                    elif header_idx is None and date_val is None and t in ("date:", "date"):
                        for k in range(j + 1, min(j + 6, len(row))):
                            d = norm_date(row[k])
                            if d:
                                date_val = d
                                break
                first_cells = [_cell_text(c).lower() for c in row[:6]]
                if header_idx is None and "adviser" in first_cells:
                    header_idx = i
                    header_row = row
            if header_idx is None:
                continue  # not the statement sheet

            colmap = {}
            for j, c in enumerate(header_row):
                t = _cell_text(c).lower()
                if t in _HEADER_MAP and _HEADER_MAP[t] not in colmap:
                    colmap[_HEADER_MAP[t]] = j
            if "adviser" not in colmap or "amount" not in colmap:
                continue

            items = []
            for row in rows[header_idx + 1:]:
                adv = _cell_text(row[colmap["adviser"]]) if colmap["adviser"] < len(row) else ""
                if not adv:
                    continue
                if adv.lower().startswith("total"):
                    break

                def g(key):
                    idx = colmap.get(key)
                    return row[idx] if idx is not None and idx < len(row) else None

                amt = norm_amount(g("amount"))
                item = {
                    "date": norm_date(g("date")),
                    "lender": _cell_text(g("lender")),
                    "ref": _cell_text(g("ref")),
                    "type": _cell_text(g("type")),
                    "firstName": _cell_text(g("firstName")),
                    "surname": _cell_text(g("surname")),
                    "class": _cell_text(g("class")),
                    "amount": amt,
                }
                if amt is None and not item["ref"] and not item["surname"]:
                    continue
                if amt is None:
                    item["amount"] = 0.0
                items.append(item)

            parsed = {"date": date_val, "total": total_val, "items": items}
            break
    finally:
        wb.close()

    if parsed is None:
        reason = "layout not matched (no Adviser table found)"
        if "unallocated" in filename.lower():
            reason = "unallocated/incomplete pipeline list, not a statement"
        return {"file": filename, "skipped": reason}

    if not parsed["date"]:
        parsed["date"] = date_from_filename(filename)

    items_sum = round(sum(i["amount"] for i in parsed["items"]), 2)
    total_derived = False
    if parsed["total"] is None:
        # Some files store the Total as an uncached =SUM() formula (no value);
        # derive it from the line items.
        parsed["total"] = items_sum
        total_derived = True

    statement = {
        "file": filename,
        "date": parsed["date"],
        "total": round(parsed["total"], 2),
        "items": parsed["items"],
    }
    if total_derived:
        statement["totalDerivedFromItems"] = True
    elif abs(items_sum - statement["total"]) > 1.0:
        statement["sumMismatch"] = True
        statement["itemsSum"] = items_sum
    return statement


if __name__ == "__main__":
    import json
    import sys
    if len(sys.argv) < 2:
        sys.exit("usage: python3 scripts/statement_parser.py <statement.xlsx> [...]")
    for p in sys.argv[1:]:
        print(json.dumps(parse_statement(p), indent=1))
