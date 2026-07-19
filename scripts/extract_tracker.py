#!/usr/bin/env python3
"""Extract the Mortgage Oasis "Tracker" Google Sheet into data/tracker.json.

Usage:
    python3 scripts/extract_tracker.py <path-to-Tracker.xlsx> [output.json]

The Tracker workbook is the Google Sheet exported as .xlsx
(File > Download > Microsoft Excel, or via the Drive API export).
It contains one "Business Register" tab per year; the column layout
changed slightly between years, so each year has its own column map.
"""
import json
import re
import sys
from datetime import datetime, date

import openpyxl

# Column maps: unified field -> 0-based column index per register tab.
LAYOUTS = {
    "Business Register 2026": {
        "first_data_row": 3,
        "cols": {
            "date": 0, "client": 1, "admin": 2, "source": 3, "property": 4,
            "business": 5, "product": 6, "provider": 7, "amount": 8,
            "premium": 9, "brokerFee": 10, "feeCollected": 11,
            "commissionWritten": 12, "completed": 19, "commissionReceived": 20,
        },
    },
    "Business Register 2025": {
        "first_data_row": 4,
        "cols": {
            "date": 0, "client": 1, "admin": 2, "source": 3, "property": 4,
            "business": 5, "product": 6, "provider": 7, "amount": 8,
            "premium": 9, "brokerFee": 10, "feeCollected": 11,
            "commissionWritten": 12, "completed": 17, "commissionReceived": 18,
        },
    },
    "Business Register 2024": {
        "first_data_row": 4,
        "cols": {
            "date": 0, "client": 1, "admin": 2, "source": 3, "property": 4,
            "business": 5, "product": None, "provider": 6, "amount": 7,
            "premium": 8, "brokerFee": 9, "feeCollected": 10,
            "commissionWritten": 11, "completed": 17, "commissionReceived": 18,
        },
    },
}

# Normalise the free-text "Type of Business" into a small set of segments
# the dashboard can chart. Raw values are kept alongside.
SEGMENT_RULES = [
    ("MCR", ["mcr"]),
    ("Protection", ["life", "cic", "critical", "income protection", "asu",
                    "protection", "whole of life", "family income"]),
    ("General Insurance", ["building", "contents", "b&c", "gi", "home insurance"]),
    ("Mortgage", ["resi", "buy to let", "buy-to-let", "btl", "mortgage",
                  "mogage", "motgage", "mtg", "remo", "purchase", "ftb",
                  "first time", "mover", "product transfer", "switch",
                  "porting", "further advance", "2nd charge", "second charge",
                  "self build", "capital raising", "bridging"]),
]


def segment_for(business, product):
    text = f"{business or ''} {product or ''}".lower()
    if not text.strip():
        return "Other"
    for segment, keywords in SEGMENT_RULES:
        if any(k in text for k in keywords):
            return segment
    return "Other"


def as_money(value):
    if value is None:
        return None
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 2)
    text = str(value).strip()
    if text in ("", "-", "TBC", "N/A", "n/a"):
        return None
    match = re.search(r"-?[\d,]+(?:\.\d+)?", text.replace("£", ""))
    if not match:
        return None
    negative = text.lstrip().startswith("-") or "(" in text
    number = float(match.group(0).replace(",", ""))
    return round(-abs(number) if negative else number, 2)


def as_iso_date(value):
    if value is None:
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%d-%m-%Y", "%d %b %Y", "%d-%b-%y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def as_text(value):
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def extract(workbook_path):
    wb = openpyxl.load_workbook(workbook_path, data_only=True)
    cases = []
    for sheet_name, layout in LAYOUTS.items():
        if sheet_name not in wb.sheetnames:
            print(f"WARNING: sheet '{sheet_name}' not found, skipping")
            continue
        ws = wb[sheet_name]
        year = int(sheet_name.rsplit(" ", 1)[1])
        cols = layout["cols"]
        for row in ws.iter_rows(min_row=layout["first_data_row"],
                                values_only=True):
            def cell(field):
                idx = cols.get(field)
                return row[idx] if idx is not None and idx < len(row) else None

            client = as_text(cell("client"))
            case_date = as_iso_date(cell("date"))
            if not client and not case_date:
                continue  # blank filler row
            business = as_text(cell("business"))
            product = as_text(cell("product"))
            cases.append({
                "register": year,
                "date": case_date,
                "client": client,
                "admin": as_text(cell("admin")),
                "source": as_text(cell("source")),
                "property": as_text(cell("property")),
                "business": business,
                "product": product,
                "segment": segment_for(business, product),
                "provider": as_text(cell("provider")),
                "amount": as_money(cell("amount")),
                "premium": as_money(cell("premium")),
                "brokerFee": as_money(cell("brokerFee")),
                "feeCollectedDate": as_iso_date(cell("feeCollected")),
                "commissionWritten": as_money(cell("commissionWritten")),
                "completedDate": as_iso_date(cell("completed")),
                "commissionReceived": as_money(cell("commissionReceived")),
            })
    return cases


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    out_path = sys.argv[2] if len(sys.argv) > 2 else "data/tracker.json"
    cases = extract(sys.argv[1])
    payload = {
        "generatedAt": datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ"),
        "source": "Tracker (Google Sheets) - Business Register tabs",
        "cases": cases,
    }
    with open(out_path, "w") as fh:
        json.dump(payload, fh, indent=1)
    written = sum(c["commissionWritten"] or 0 for c in cases)
    received = sum(c["commissionReceived"] or 0 for c in cases)
    print(f"{len(cases)} cases -> {out_path}")
    print(f"commission written £{written:,.2f} | received £{received:,.2f}")
    for year in sorted({c['register'] for c in cases}):
        n = sum(1 for c in cases if c["register"] == year)
        print(f"  {year}: {n} cases")


if __name__ == "__main__":
    main()
