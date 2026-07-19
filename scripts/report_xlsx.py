#!/usr/bin/env python3
"""Format the dashboard's report data as a styled Excel workbook.

The dashboard POSTs its own computed report rows (the same numbers it shows,
including statement matching and manual overrides) to /api/report; this module
only does the formatting, so there is exactly one source of truth for the
figures. Returns the .xlsx as bytes.
"""
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ACCENT = "2A78D6"
INK_2 = "52514E"
GBP = '£#,##0.00'

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor=ACCENT)
TITLE_FONT = Font(bold=True, size=14)
META_FONT = Font(color=INK_2, size=10)


def _sheet(wb, title, headers, widths):
    ws = wb.create_sheet(title)
    for i, (header, width) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    return ws


def _money(ws, cell_ref):
    ws[cell_ref].number_format = GBP


def build_report(data):
    wb = Workbook()

    # ---- Summary ----
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Mortgage Oasis — Commission report"
    ws["A1"].font = TITLE_FONT
    ws["A2"] = (f"Period: {data.get('period', 'All years')}  ·  Generated "
                f"{str(data.get('generatedAt', ''))[:10]}  ·  Overdue after "
                f"{data.get('thresholds', {}).get('overdueDays', '?')} days  ·  "
                f"Variance threshold {data.get('thresholds', {}).get('variancePct', '?')}%")
    ws["A2"].font = META_FONT
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 18
    row = 4
    for label, value in data.get("summary", []):
        ws.cell(row=row, column=1, value=label)
        cell = ws.cell(row=row, column=2, value=value)
        if isinstance(value, (int, float)) and "ases" not in str(label):
            cell.number_format = GBP
        cell.alignment = Alignment(horizontal="right")
        row += 1

    # ---- Outstanding ----
    ws = _sheet(wb, "Outstanding",
                ["Date", "Client", "Business", "Provider", "Predicted",
                 "Days waiting", "Status", "Chased on", "Follow-up due"],
                [11, 32, 18, 18, 13, 12, 14, 11, 12])
    for r, item in enumerate(data.get("outstanding", []), start=2):
        ws.cell(row=r, column=1, value=item.get("date"))
        ws.cell(row=r, column=2, value=item.get("client"))
        ws.cell(row=r, column=3, value=item.get("business"))
        ws.cell(row=r, column=4, value=item.get("provider"))
        ws.cell(row=r, column=5, value=item.get("predicted"))
        _money(ws, f"E{r}")
        ws.cell(row=r, column=6, value=item.get("days"))
        status = ws.cell(row=r, column=7, value=item.get("status"))
        if item.get("status") in ("Overdue", "Follow up due"):
            status.font = Font(bold=True, color="D03B3B")
        ws.cell(row=r, column=8, value=item.get("chasedOn"))
        ws.cell(row=r, column=9, value=item.get("followUpDue"))

    # ---- Variance ----
    ws = _sheet(wb, "Paid vs predicted",
                ["Date", "Client", "Provider", "Predicted", "Paid",
                 "Shortfall", "Variance %", "Paid via"],
                [11, 32, 18, 13, 13, 13, 11, 15])
    for r, item in enumerate(data.get("variance", []), start=2):
        ws.cell(row=r, column=1, value=item.get("date"))
        ws.cell(row=r, column=2, value=item.get("client"))
        ws.cell(row=r, column=3, value=item.get("provider"))
        ws.cell(row=r, column=4, value=item.get("predicted"))
        ws.cell(row=r, column=5, value=item.get("paid"))
        ws.cell(row=r, column=6, value=item.get("shortfall"))
        for col in ("D", "E", "F"):
            _money(ws, f"{col}{r}")
        ws.cell(row=r, column=7, value=item.get("variancePct"))
        ws.cell(row=r, column=8, value=item.get("via"))

    # ---- Monthly income ----
    ws = _sheet(wb, "Monthly income",
                ["Month", "Statement income", "Non-indemnity (NI)", "Recurring (R)"],
                [11, 17, 17, 14])
    for r, item in enumerate(data.get("monthly", []), start=2):
        ws.cell(row=r, column=1, value=item.get("month"))
        ws.cell(row=r, column=2, value=item.get("income"))
        ws.cell(row=r, column=3, value=item.get("ni"))
        ws.cell(row=r, column=4, value=item.get("r"))
        for col in ("B", "C", "D"):
            _money(ws, f"{col}{r}")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
